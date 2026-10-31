import hashlib
import math
import re
from datetime import UTC, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from pathlib import Path
from typing import Literal, get_args

from openpyxl import load_workbook
from pydantic import BaseModel, Field

REFERENCE = Path(__file__).resolve().parent.parent / "data" / "reference"
SUFFIXES = {"catalogue": ".xlsx", "quote-template": ".docx"}
ITEM_FIELDS = ("code", "description", "category", "unit", "unit_price", "max_quantity")
MEASURED = ("m", "m2", "l", "h")
MIN_CONFIDENCE = 0.6
CENT = Decimal("0.01")

HeaderField = Literal["lift_id", "location", "inspection_date", "technician"]
HEADER_FIELDS: tuple[HeaderField, ...] = get_args(HeaderField)
Flag = Literal["missing", "low_confidence", "ambiguous", "no_match", "invalid_code", "invalid_quantity"]
Outcome = Literal["proposed", "accepted", "changed", "removed", "added"]


class DomainError(ValueError):
    """A change that the quote's rules refuse; the message is shown to the estimator."""


class Version(BaseModel):
    name: str
    sha256: str


class Item(BaseModel):
    code: str
    description: str
    category: str
    unit: str
    unit_price: Decimal
    max_quantity: Decimal


class Catalogue(BaseModel):
    version: Version
    vat_rate: Decimal
    items: dict[str, Item]


class Defect(BaseModel):
    page: int
    row: int
    text: str


class Match(BaseModel):
    expanded: str
    category: str
    item_code: str | None
    quantity: float
    confidence: float
    candidates: list[str]
    reason: str | None


class Line(BaseModel):
    number: int
    defect: Defect | None = None
    expanded: str = ""
    item_code: str | None = None
    description: str = ""
    unit: str = ""
    quantity: Decimal = Decimal(1)
    unit_price: Decimal | None = None
    amount: Decimal | None = None
    candidates: list[str] = []
    flag: Flag | None = None
    reason: str = ""
    outcome: Outcome = "proposed"


class Header(BaseModel):
    lift_id: str | None = None
    location: str | None = None
    inspection_date: str | None = None
    technician: str | None = None


class Event(BaseModel):
    at: datetime
    by: str
    action: str
    detail: str


class Edit(BaseModel):
    action: Literal["accept", "item", "quantity", "remove", "add", "header"]
    line: int | None = None
    code: str | None = None
    quantity: float | None = None
    field: HeaderField | None = None
    value: str | None = None


def now() -> datetime:
    return datetime.now(UTC)


class Quote(BaseModel):
    id: str
    state: Literal["drafting", "review", "approved", "failed"] = "drafting"
    catalogue: Version
    template: Version
    header: Header = Field(default_factory=Header)
    header_flags: dict[HeaderField, Flag] = {}
    lines: list[Line] = []
    vat_rate: Decimal = Decimal(0)
    subtotal: Decimal = Decimal(0)
    vat: Decimal = Decimal(0)
    total: Decimal = Decimal(0)
    events: list[Event] = []
    created_at: datetime = Field(default_factory=now)
    drafted_at: datetime | None = None
    approved_at: datetime | None = None
    approved_by: str | None = None


def money(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def reference(kind: str, pinned: Version | None = None, folder: Path = REFERENCE) -> tuple[bytes, Version]:
    suffix = SUFFIXES[kind]
    pattern = re.compile(rf"{kind}-v(\d+){re.escape(suffix)}")
    found = sorted((int(match[1]), path) for path in folder.iterdir() if (match := pattern.fullmatch(path.name)))
    if pinned is None and not found:
        raise DomainError(f"the reference data has no {kind}")
    path = folder / f"{pinned.name}{suffix}" if pinned else found[-1][1]
    if not path.is_file():
        raise DomainError(f"{path.name} is not in the reference data")
    data = path.read_bytes()
    version = Version(name=path.stem, sha256=hashlib.sha256(data).hexdigest())
    if pinned is not None and version != pinned:
        raise DomainError(f"{pinned.name} changed after this quote used it; release a new version instead")
    return data, version


def load_catalogue(pinned: Version | None = None, folder: Path = REFERENCE) -> Catalogue:
    data, version = reference("catalogue", pinned, folder)
    workbook = load_workbook(BytesIO(data), data_only=True)
    rows = workbook["items"].iter_rows(min_row=2, values_only=True)
    items = [Item.model_validate(dict(zip(ITEM_FIELDS, map(str, row), strict=True))) for row in rows if row[0]]
    vat_rate = Decimal(str(workbook["terms"]["B1"].value))
    return Catalogue(version=version, vat_rate=vat_rate, items={item.code: item for item in items})


def new_quote(quote_id: str, folder: Path = REFERENCE) -> Quote:
    catalogue, template = reference("catalogue", None, folder)[1], reference("quote-template", None, folder)[1]
    return Quote(id=quote_id, catalogue=catalogue, template=template)


def propose(number: int, defect: Defect, match: Match, catalogue: Catalogue) -> Line:
    candidates = list(dict.fromkeys(code for code in match.candidates if code in catalogue.items))[:3]
    quantity = Decimal(str(match.quantity)) if math.isfinite(match.quantity) else Decimal(0)
    line = Line(number=number, defect=defect, expanded=match.expanded, candidates=candidates, quantity=quantity)
    line.reason = match.reason or ""
    code = match.item_code
    if code is None:
        line.flag = "ambiguous" if candidates else "no_match"
    elif code not in catalogue.items or (match.candidates and code not in candidates):
        line.flag, line.reason = "invalid_code", f"{code[:40]} is not a valid choice in {catalogue.version.name}"
    else:
        line.item_code = code
        line.flag = None if match.confidence >= MIN_CONFIDENCE else "low_confidence"
    return line


def valid(quantity: Decimal, item: Item) -> bool:
    if not 0 < quantity <= item.max_quantity:
        return False
    return quantity == money(quantity) and (item.unit in MEASURED or quantity % 1 == 0)


def price(quote: Quote, catalogue: Catalogue) -> Quote:
    if catalogue.version != quote.catalogue:
        raise DomainError(f"this quote is priced with {quote.catalogue.name}, not {catalogue.version.name}")
    for line in quote.lines:
        item = catalogue.items.get(line.item_code or "") if line.outcome != "removed" else None
        line.flag = None if line.flag == "invalid_quantity" else line.flag
        line.unit_price = line.amount = None
        if item is None:
            continue
        line.description, line.unit = item.description, item.unit
        if not valid(line.quantity, item):
            line.flag = line.flag or "invalid_quantity"
            continue
        line.unit_price, line.amount = item.unit_price, money(line.quantity * item.unit_price)
    quote.vat_rate = catalogue.vat_rate
    quote.subtotal = money(sum((line.amount for line in quote.lines if line.amount is not None), Decimal(0)))
    quote.vat = money(quote.subtotal * quote.vat_rate)
    quote.total = quote.subtotal + quote.vat
    return quote


def draft(quote: Quote, proposals: list[tuple[Defect, Match]], catalogue: Catalogue) -> Quote:
    if quote.state != "drafting":
        raise DomainError(f"the quote is {quote.state} and already has a draft")
    drafted = quote.model_copy(deep=True)
    drafted.lines = [propose(number, defect, match, catalogue) for number, (defect, match) in enumerate(proposals, 1)]
    drafted.state, drafted.drafted_at = "review", now()
    return price(drafted, catalogue)


def open_flags(quote: Quote) -> list[str]:
    kept = [line for line in quote.lines if line.outcome != "removed"]
    flags = [f"{name}: missing" for name in HEADER_FIELDS if not (getattr(quote.header, name) or "").strip()]
    flags += [f"{name}: {flag}" for name, flag in quote.header_flags.items()]
    flags += [f"line {line.number}: {line.flag}" for line in kept if line.flag]
    return flags if kept else [*flags, "the quote has no lines"]


def editable(quote: Quote) -> Quote:
    if quote.state != "review":
        raise DomainError(f"the quote is {quote.state} and can no longer change")
    return quote.model_copy(deep=True)


def find(quote: Quote, number: int | None) -> Line:
    line = next((line for line in quote.lines if line.number == number and line.outcome != "removed"), None)
    if line is None:
        raise DomainError(f"line {number} is not in this quote")
    return line


def known(code: str | None, catalogue: Catalogue) -> str:
    if code is None or code not in catalogue.items:
        raise DomainError(f"{code} is not in {catalogue.version.name}")
    return code


def units(quantity: float | None) -> Decimal:
    if quantity is None or not math.isfinite(quantity) or quantity <= 0:
        raise DomainError("a quantity must be a number greater than zero")
    return Decimal(str(quantity))


def apply(quote: Quote, edit: Edit, catalogue: Catalogue) -> None:
    if edit.action == "header":
        value = (edit.value or "").strip()
        if edit.field is None or not value:
            raise DomainError(f"a header edit needs a value for one of {', '.join(HEADER_FIELDS)}")
        setattr(quote.header, edit.field, value)
        quote.header_flags.pop(edit.field, None)
        return
    if edit.action == "add":
        code, quantity = known(edit.code, catalogue), units(edit.quantity)
        quote.lines.append(Line(number=len(quote.lines) + 1, item_code=code, quantity=quantity, outcome="added"))
        return
    line = find(quote, edit.line)
    if edit.action == "accept" and (line.item_code is None or line.flag not in (None, "low_confidence")):
        raise DomainError(f"line {line.number} needs an item and a valid quantity before it can be accepted")
    before = (line.item_code, line.quantity)
    if edit.action == "item":
        line.item_code = known(edit.code, catalogue)
    if edit.action == "quantity":
        line.quantity = units(edit.quantity)
    else:
        line.flag = None
    if edit.action == "remove":
        line.outcome = "removed"
    elif edit.action == "accept" and line.outcome == "proposed":
        line.outcome = "accepted"
    elif (line.item_code, line.quantity) != before and line.outcome != "added":
        line.outcome = "changed"


def review(quote: Quote, edits: list[Edit], by: str, catalogue: Catalogue) -> Quote:
    updated = editable(quote)
    for edit in edits:
        apply(updated, edit, catalogue)
        detail = edit.model_dump_json(exclude_none=True)
        updated.events.append(Event(at=now(), by=by, action=edit.action, detail=detail))
    return price(updated, catalogue)


def approve(quote: Quote, by: str, catalogue: Catalogue, folder: Path = REFERENCE) -> Quote:
    approved = price(editable(quote), catalogue)
    reference("quote-template", approved.template, folder)
    if flags := open_flags(approved):
        raise DomainError(f"resolve every flag before approving: {'; '.join(flags)}")
    for line in approved.lines:
        line.outcome = "accepted" if line.outcome == "proposed" else line.outcome
    approved.state, approved.approved_by, approved.approved_at = "approved", by, now()
    approved.events.append(Event(at=approved.approved_at, by=by, action="approve", detail=""))
    return approved
