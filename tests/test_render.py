from io import BytesIO

import pytest
from docx import Document

from app.domain import Catalogue, DomainError, Edit, Quote, approve, review
from app.render import render


def test_the_word_quote_traces_every_line_to_its_source(drafted: Quote, catalogue: Catalogue) -> None:
    edits = [Edit(action="item", line=3, code="ELE-105"), Edit(action="accept", line=4), Edit(action="remove", line=2)]
    drafted.header.location = f"{drafted.header.location}\x0c"
    edited = review(drafted, [*edits, Edit(action="add", code="LAB-102", quantity=1)], "estimator-1", catalogue)
    quote = approve(edited, "estimator-1", catalogue)
    document = Document(BytesIO(render(quote)))
    rows = [[cell.text for cell in row.cells] for row in document.tables[1].rows[1:]]
    assert [(row[0], row[2]) for row in rows] == [("1", "2.5 h"), ("2", "4"), ("3", "2"), ("4", "1 h")]
    sources = {"LAB-101": "report page 1, row 1", "ELE-105": "report page 2, row 1", "DOR-102": "report page 2, row 2"}
    sources["LAB-102"] = "added in review"
    assert [row[1].splitlines()[1] for row in rows] == [f"{c} · catalogue-v1 · {s}" for c, s in sources.items()]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    text = "\n".join([paragraph.text for paragraph in document.paragraphs] + cells)
    assert "6 Mill Lane & Chapel Row, Northgate" in text and f"€{quote.total:,.2f}" in text
    assert quote.catalogue.sha256[:12] in text and "quote-template-v1" in text and "{{" not in text and "{%" not in text


def test_no_document_before_approval(drafted: Quote) -> None:
    with pytest.raises(DomainError, match="only an approved quote has a document"):
        render(drafted)
