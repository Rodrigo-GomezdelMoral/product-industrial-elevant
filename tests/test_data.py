import hashlib
import json
from collections import Counter
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook

from app.domain import REFERENCE, Version, load_catalogue, money
from ops.make_reports import REPORTS, TRUTH


@pytest.mark.parametrize("path", sorted(REFERENCE.glob("catalogue-v*.xlsx")), ids=lambda path: path.name)
def test_every_catalogue_version_is_well_formed(path: Path) -> None:
    pinned = Version(name=path.stem, sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    catalogue = load_catalogue(pinned, path.parent)
    assert len(catalogue.items) == load_workbook(path)["items"].max_row - 1 and 0 <= catalogue.vat_rate < 1
    for item in catalogue.items.values():
        assert item.code == item.code.strip() and item.description.strip() not in ("", "None"), item.code
        assert 0 < item.unit_price == money(item.unit_price) and 0 < item.max_quantity == int(item.max_quantity)


def test_reports_match_their_truth() -> None:
    records = [json.loads(line) for line in TRUTH.read_text(encoding="utf-8").splitlines()]
    codes = set(load_catalogue().items)
    kinds = Counter(defect["kind"] for record in records for defect in record["defects"])
    assert (len(records), sum(record["scanned"] for record in records), kinds["injection"]) == (40, 10, 2)
    assert 0.13 < (kinds["ambiguous"] + kinds["no_match"]) / kinds.total() < 0.17
    for record in records:
        data = (REPORTS / record["report"]).read_bytes()
        pages = pdfium.PdfDocument(data)
        assert len(data) < 10 * 2**20 and len(pages) <= 4
        for defect in record["defects"]:
            assert {defect["item_code"], *defect["candidates"]} - {None} <= codes
            text = " ".join(pages[defect["page"] - 1].get_textpage().get_text_range().split())
            assert record["scanned"] or " ".join(defect["component"].split()) in text, record["report"]
